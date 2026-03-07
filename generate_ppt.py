from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title, content_lines):
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    if not content_lines:
        return
    tf.text = content_lines[0]
    
    for line in content_lines[1:]:
        p = tf.add_paragraph()
        if line.startswith("    "):
            p.text = line.strip()
            p.level = 1
        else:
            p.text = line
            p.level = 0

# Slide 1
add_slide("The Problem", [
    "• The Struggle of Fragmented Ideas: We all experience highly vivid, emotional thoughts, but they rarely arrive perfectly structured.",
    "• The Blank Page Anxiety: Sitting down to type breaks the psychological flow. The emotion fades while trying to find the 'right words.'",
    "• The Limitation of Voice Memos: Unstructured voice memos pile up and are overwhelming to listen to, re-organize, and process into something readable."
])

# Slide 2
add_slide("The Solution - Miss Writer ✨", [
    "• What it is: An end-to-end Voice-to-Story AI Agent designed to bridge the gap between spoken thought and written narrative.",
    "• The User Experience: The user simply speaks their chaotic, unorganized thoughts out loud.",
    "• The Output: The system returns a beautifully written, structured narrative that maintains the user's original emotional intent.",
    "• The 'Why': Created to allow powerful untold stories and voices to be effortlessly expressed to the world."
])

# Slide 3
add_slide("End-to-End System Architecture", [
    "1. Audio Capture (Frontend): Browser-based high-quality microphone capture using Streamlit.",
    "2. Transcription (API): Fast robust conversion using OpenAI's Whisper model over REST.",
    "3. The Brain (LangGraph): Multi-step asynchronous routing of the raw text through a graph of specialized AI Agents.",
    "4. Delivery: The final generated story is streamed back instantly to the frontend UI for live consumption."
])

# Slide 4
add_slide("The Agentic Approach", [
    "• The Traditional Wrapper Approach: Prompt = 'Turn this transcript into a cohesive story.' -> generic, robotic results.",
    "• The 'Miss Writer' LangGraph Approach:",
    "    - Node 1 (Thought Interpreter): Employs GPT-4o-mini to analyze psychological subtext and extract Themes, Emotional Tone, and Core Ideas.",
    "    - Node 2 (Story Writer): Employs GPT-4o to dynamically generate the final polished narrative using the emotional constraints.",
    "• The Result: Better intent preservation, higher quality aesthetic writing, and significantly reduced AI hallucination."
])

# Slide 5
add_slide("The Tech Stack", [
    "• LangGraph: Core framework for defining and routing stateful multi-actor agent flows.",
    "• FastAPI: Highly performant, async Python framework providing the core backend microservice.",
    "• Streamlit: Rapid frontend GUI with custom CSS for a premium user experience.",
    "• Docker & Docker-Compose: Full containerization for consistent deployment.",
    "• Railway.app: Instant, serverless cloud deployment pipeline for Dockerized applications."
])

# Slide 6
add_slide("Future Roadmap & Extensibility", [
    "• Persistent Storage: Hooking FastAPI to a Postgres or MongoDB database to save user histories.",
    "• Voice Cloning & TTS: Passing the generated story back into a TTS pipeline so the app speaks the finalized story back to the user.",
    "• Multi-Platform Output: Adding LangGraph nodes to format the story simultaneously into a 'LinkedIn Post', 'Twitter Thread', and 'Medium Article'."
])

prs.save('/Users/sandhyabantiduttaborah/Desktop/Miss_Writer_Presentation.pptx')
print("Successfully generated Miss_Writer_Presentation.pptx on the Desktop!")
